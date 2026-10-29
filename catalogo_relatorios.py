"""Monta o catálogo de relatórios e exportações do SPICA 7 e do SPICA 8 a partir das pastas do levantamento 6699.

Cada rotina (tela) vira uma linha: o código do arquivo (x010, w745 na V8; l300, p100 na V7) identifica a rotina e os
arquivos com o mesmo código são os layouts dela. As planilhas geradas servem para revisar os nomes antes de alimentar o buscador.
"""
import html
import os
import re
import sqlite3
import unicodedata
import zipfile
from collections import defaultdict
from pathlib import Path

import pandas as pd
import pypdf
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from rapidfuzz import fuzz

from relatorios import limpar_titulo

try:  # PyMuPDF separa as palavras do PDF melhor que o pypdf ("StatusICMS" x "Status ICMS")
    import pymupdf
except ImportError:  # versões antigas
    import fitz as pymupdf

RAIZ = Path(os.environ.get("RELATORIOS_DIR", r"C:\6699_Levantamento_Inicial"))
PLANILHA = RAIZ / "6699_Levantamento_Inicial_Lista.ods"
BANCO = Path(__file__).resolve().parent / "relatorios.db"  # o que o !relatorio do bot consulta
VERSOES = ("V8", "V7")
# Em cada versão: "Vx" = só os que citam imposto (os da planilha 6699); "Vx TODOS" = tudo o que foi exportado
PASTAS = {"V8": ("V8 TODOS", "V8"), "V7": ("V7 TODOS", "V7")}
ABA_PLANILHA = {"V8": "SPICA 8", "V7": "SPICA7"}
LIMITE_TEXTO = 100_000  # caracteres guardados por arquivo


def saida(versao):
    return Path(__file__).resolve().parent / f"relatorios_{versao.lower()}_rotinas.xlsx"

# Nome proposto de cada rotina e a área onde ela fica (palpite pelo código e pelo diário do levantamento)
ROTINAS = {  # código -> (nome da tela no menu, aba > grupo do menu de Relatórios)
    "u160": ("Produtos com Filtros", "R.Básicos > Produtos"),
    "u240": ("Produtos Por Tributação", "R.Básicos > Produtos"),
    "u260": ("Produtos Por NCM", "R.Básicos > Produtos"),
    "u480": ("Produtos - TUSS", "R.Básicos > Produtos"),
    "manual-cadastro-produto": ("Exporta Cadastro p/ Excel", "R.Básicos > Produtos"),
    "manual-cadastro-bionexo": ("Exporta Produtos para Bionexo", "R.Básicos > Produtos"),
    "v040": ("Posição do Estoque (Fechamento)", "R.Estoque > Posição do Estoque"),
    "manual-posicao-estoque": ("Posição do Estoque - Exportação (Fechamento)", "R.Estoque > Posição do Estoque"),
    "v620": ("Notas Fiscais Recebidas", "R.Estoque > Compras / Entrada de Mercadorias"),
    "manual-nf-compra": ("Exportação de Notas Fiscais de Compra", "R.Estoque > Compras / Entrada de Mercadorias"),
    "v670": ("Exportação de Notas Fiscais Recebidas", "R.Estoque > Compras / Entrada de Mercadorias"),
    "w100": ("Cirurgias Exportação p/Excel", "R.Gerenciais > Cirurgias"),
    "w450": ("Faturamento - Análise de Lucro Bruto", "R.Gerenciais > Análise de Lucratividade"),
    "w470": ("Análise de Lucratividade por NF", "R.Gerenciais > Análise de Lucratividade"),
    "w472": ("Análise de Lucratividade de NF de Venda (Tipo II)", "R.Gerenciais > Análise de Lucratividade"),
    "w490": ("Apuração de Lucro por NF", "R.Gerenciais > Análise de Lucratividade"),
    "w620": ("Faturamento", "R.Gerenciais II > Faturamento e Analises"),
    "w640": ("Faturamento Por Vendedor", "R.Gerenciais II > Faturamento e Analises"),
    "w655": ("Faturamento Resumido", "R.Gerenciais II > Faturamento e Analises"),
    "w670": ("Faturamento Analítico", "R.Gerenciais II > Faturamento e Analises"),
    "w710": ("Notas Fiscais Emitidas por Faturista", "R.Gerenciais II > Faturamento e Analises"),
    "w725": ("Faturamento Mensal / Anual por Segmento e por Vendedor", "R.Gerenciais II > Diversos / Exportação"),
    "w730": ("Relatório Anual de Vendas por Vendedor / Cliente", "R.Gerenciais II > Diversos / Exportação"),
    "w735": ("Gráficos - Faturamento", "R.Gerenciais II > Diversos / Exportação"),
    "w740": ("Notas Fiscais Canceladas", "R.Gerenciais II > Diversos / Exportação"),
    "w745": ("Exportação de Notas Fiscais de Venda", "R.Gerenciais II > Diversos / Exportação"),
    "w750": ("Relatório de Vendas por Vendedor - Exportação", "R.Gerenciais II > Diversos / Exportação"),
    "w755": ("Exportação de NF Emitidas e Recebidas", "R.Gerenciais II > Diversos / Exportação"),
    "w860": ("Exportação de Notas Fiscais (Modelo II)", "R.Gerenciais II > Diversos / Exportação"),
    "w862": ("Exportação de Notas Fiscais (Modelo III)", "R.Gerenciais II > Diversos / Exportação"),
    "w865": ("Exportação de NF Emitidas", "R.Gerenciais II > Diversos / Exportação"),
    "w867": ("Exportação - Mensal de Vendas II", "R.Gerenciais II > Diversos / Exportação"),
    "w770": ("Curva ABC de Vendas por Cliente", "R.Gerenciais II > Curva ABC"),
    "w775": ("Curva ABC de Vendas por Produto", "R.Gerenciais II > Curva ABC"),
    "w785": ("Curva ABC de Vendas por Fabricante", "R.Gerenciais II > Curva ABC"),
    "w790-isentos": ("Vendas de Produtos Isentos PIS/COFINS", "R.Gerenciais II > PIS/COFINS"),
    "w790-tributados": ("Vendas de Produtos Tributados PIS/COFINS", "R.Gerenciais II > PIS/COFINS"),
    "w790-apuracao": ("Apuração de PIS/COFINS", "R.Gerenciais II > PIS/COFINS"),
    "w790-analitica": ("Apuração de PIS/COFINS Analítica", "R.Gerenciais II > PIS/COFINS"),
    "w790-recebimento": ("Apuração de PIS/COFINS - Data Recebimento", "R.Gerenciais II > PIS/COFINS"),
    "w790-nf-recebidas": ("Apuração de PIS/COFINS Analítica - NF Recebidas", "R.Gerenciais II > PIS/COFINS"),
    "w790-natureza": ("Apuração de PIS/COFINS - Com Natureza da Rec.", "R.Gerenciais II > PIS/COFINS"),
    "w820": ("Apuração de ICMS e IPI", "R.Gerenciais II > ICMS e IPI"),
    "x010": ("Contas a Pagar Por Vencimento", "R.Financ. > Contas a Pagar"),
    "x020": ("Contas a Pagar por Emissão", "R.Financ. > Contas a Pagar"),
    "x030": ("Contas a Pagar por Fornecedor", "R.Financ. > Contas a Pagar"),
    "x040": ("Contas a Pagar por Conta do Plano", "R.Financ. > Contas a Pagar"),
    "x050": ("Contas a Pagar - Exportação p/ Excel", "R.Financ. > Contas a Pagar"),
    "x055": ("Contas a Pagar - Visão em Data Retroativa", "R.Financ. > Contas a Pagar"),
    "x060": ("Contas Pagas", "R.Financ. > Contas Pagas"),
    "x070": ("Contas Pagas Por Fornecedor", "R.Financ. > Contas Pagas"),
    "x080": ("Contas Pagas Por Forma", "R.Financ. > Contas Pagas"),
    "x090": ("Contas Pagas por Fornecedor / Emissão", "R.Financ. > Contas Pagas"),
    "x100": ("Contas a Pagar / Pagas Por Forma", "R.Financ. > Auxiliares de Despesas"),
    "x102": ("Contas a Pagar / Pagas", "R.Financ. > Auxiliares de Despesas"),
    "x110": ("Despesas por Conta do Plano", "R.Financ. > Auxiliares de Despesas"),
    "x120": ("Despesas por Conta do Plano Mês a Mês", "R.Financ. > Auxiliares de Despesas"),
    "x130": ("Extrato de Antecipação para Fornecedor", "R.Financ. > Auxiliares de Despesas"),
    "x140": ("Comissões", "R.Financ. > Auxiliares de Despesas"),
    "x150": ("Comissões com Itens", "R.Financ. > Auxiliares de Despesas"),
    "x170": ("Relatório Mensal de Pagamento de Comissão", "R.Financ. > Auxiliares de Despesas"),
    "x180": ("Contas a Receber por Vencimento", "R.Financ. > Contas a Receber"),
    "x190": ("Contas a Receber por Emissão", "R.Financ. > Contas a Receber"),
    "x200": ("Contas a Receber por Vendedor", "R.Financ. > Contas a Receber"),
    "x220": ("Contas a Receber por Cliente", "R.Financ. > Contas a Receber"),
    "x230": ("Contas a Receber em Atraso", "R.Financ. > Contas a Receber"),
    "x255": ("Contas a Receber - Visão em Data Retroativa", "R.Financ. > Contas a Receber"),
    "x340": ("Contas a Receber / Recebidas por No.", "R.Financ. > Auxiliares de Receitas"),
    "x350": ("Ficha de Duplicatas", "R.Financ. > Auxiliares de Receitas"),
    "x360": ("Clientes em Atraso Mês a Mês", "R.Financ. > Auxiliares de Receitas"),
    "x370": ("Curva ABC Inadimplência", "R.Financ. > Auxiliares de Receitas"),
    "x380": ("Relatório de Inadimplência", "R.Financ. > Auxiliares de Receitas"),
    "x390": ("Extrato de Antecipação de Cliente", "R.Financ. > Auxiliares de Receitas"),
    "x260": ("Contas Recebidas", "R.Financ. > Contas Recebidas"),
    "x270": ("Contas Recebidas por Cliente", "R.Financ. > Contas Recebidas"),
    "x280": ("Contas Recebidas por Forma", "R.Financ. > Contas Recebidas"),
    "x290": ("Contas Recebidas Por NF", "R.Financ. > Contas Recebidas"),
    "x300": ("Relatório de Contas Recebidas Por NF Liquidada", "R.Financ. > Contas Recebidas"),
    "x450": ("Mapa de Despesas x Receitas - 1 Mês", "R.Financ. > Mapa"),
    "x480": ("Mapa de Despesas x Receitas", "R.Financ. > Mapa"),
    "x490": ("Mapa de Despesas x Receitas por Competencia", "R.Financ. > Mapa"),
    "x510": ("Mapa de Despesas x Receitas - Comparativo", "R.Financ. > Mapa"),
    "x520": ("Mapa de Previsto x Realizado - Comparativo", "R.Financ. > Mapa"),
    "x400": ("Extrato de Conta Corrente", "R.Financ. > Extrato"),
    "x410": ("Extrato Consolidado", "R.Financ. > Extrato"),
    "x245": ("Exportação de Doc. em Aberto por Plano de C.", "R.Financ. > Auxiliares"),
    "x415": ("Exportação de Débitos / Créditos / CP / CR", "R.Financ. > Auxiliares"),
    "x420": ("Fluxo de Caixa", "R.Financ. > Auxiliares"),
    "x430": ("Balanço Geral", "R.Financ. > Auxiliares"),
    "x440": ("Lançamentos por Conta do Plano", "R.Financ. > Auxiliares"),
    "x470": ("Lançamentos Sem Conta do Plano", "R.Financ. > Auxiliares"),
    "x475": ("DRE", "R.Financ. > Auxiliares"),
}

# Telas que entraram quando a V8 foi refeita com o menu inteiro (prints do menu de 29/10/2026). As exportações sem
# código no nome (ods-...) foram ligadas ao item do menu pelo nome do arquivo.
_B, _E, _G, _G2, _F, _A = "R.Básicos", "R.Estoque", "R.Gerenciais", "R.Gerenciais II", "R.Financ.", "R. Auxiliares"
MENU_V8 = {
    "u040": ("Produtos Por Tipo", f"{_B} > Produtos"), "u060": ("Produtos por Fornecedor", f"{_B} > Produtos"),
    "u080": ("Produtos Com Controle de Lote", f"{_B} > Produtos"), "u100": ("Produtos Sem Controle de Lote", f"{_B} > Produtos"),
    "u120": ("Produtos Sem Registro na ANVISA", f"{_B} > Produtos"), "u140": ("Produtos Por Registro na ANVISA", f"{_B} > Produtos"),
    "u180": ("Produtos Sem Foto", f"{_B} > Produtos"), "u200": ("Produtos Por Família", f"{_B} > Produtos"),
    "u220": ("Produtos Sem Família", f"{_B} > Produtos"),
    "u300": ("Tabela de Preços", f"{_B} > Tabelas de Preço"), "ods-tabela-precos": ("Exportação de Tabela de Preços", f"{_B} > Tabelas de Preço"),
    "u340": ("Caixas", f"{_B} > Caixas"), "u360": ("Caixas Vendidas", f"{_B} > Caixas"), "u380": ("Caixas com Valor", f"{_B} > Caixas"),
    "u400": ("Documentos a Vencer por Fornecedor", f"{_B} > Outras"), "u420": ("Documentos a Vencer por Cliente", f"{_B} > Outras"),
    "u440": ("Registros na ANVISA a Vencer", f"{_B} > Outras"), "u460": ("Clientes Bloqueados", f"{_B} > Outras"),

    "v020": ("Posição do Estoque", f"{_E} > Posição do Estoque"), "v060": ("Posição do Estoque Com Previsões", f"{_E} > Posição do Estoque"),
    "v080": ("Posição do Estoque Com Valorização", f"{_E} > Posição do Estoque"),
    "v100": ("Posição do Estoque Que Está Fora", f"{_E} > Posição do Estoque"),
    "v120": ("Posição do Estoque Com Indicador de Propriedade", f"{_E} > Posição do Estoque"),
    "v140": ("Inventário Valorizado", f"{_E} > Posição do Estoque"),
    "ods-pos-estoque": ("Posição do Estoque - Exportação", f"{_E} > Posição do Estoque"),
    "ods-pos-estoque-almox": ("Posição do Estoque - Exportação (Detalhada por Almoxarifado)", f"{_E} > Posição do Estoque"),
    "ods-pos-estoque-consig": ("Posição do Estoque - Exportação (Detalhada por Consignação)", f"{_E} > Posição do Estoque"),
    "v200": ("Posição do Estoque - Ativo Fixo", f"{_E} > Posição do Estoque"), "v180": ("Posição do Estoque por RVS", f"{_E} > Posição do Estoque"),
    "v320": ("Movimentação do Estoque", f"{_E} > Movimentações"), "ods-mov-estoque": ("Movimentação do Estoque - Exportação", f"{_E} > Movimentações"),
    "v360": ("Etiquetas Impressas (Movimentação de Etiquetas)", f"{_E} > Movimentações"),
    "v340": ("Movimentação de Estoque de Produtos - Kardex", f"{_E} > Movimentações"),
    "v220": ("Produtos em Falta", f"{_E} > Em Falta"), "v240": ("Produtos em Falta nas Caixas", f"{_E} > Em Falta"),
    "v260": ("Produtos em Falta em Pedidos", f"{_E} > Em Falta"), "v300": ("Pendências com Estoque Disponível", f"{_E} > Em Falta"),
    "v680": ("Produtos com Lotes a Vencer", f"{_E} > Auxiliares"), "v720": ("Produtos em Quarentena", f"{_E} > Auxiliares"),
    "v780": ("Solicitações para o Almoxarifado", f"{_E} > Auxiliares"), "v730": ("Relatório de Não Conformes", f"{_E} > Auxiliares"),
    "ods-rastreabilidade": ("Rastreabilidade (Exportação)", f"{_E} > Auxiliares"),
    "v790": ("Solicitações Pendentes no Almoxarifado", f"{_E} > Auxiliares"),
    "v380": ("Requisições", f"{_E} > Compras / Entrada de Mercadorias"), "v400": ("Pedidos de Compra", f"{_E} > Compras / Entrada de Mercadorias"),
    "v420": ("Compras Sintética", f"{_E} > Compras / Entrada de Mercadorias"), "v480": ("Compras Mês a Mês", f"{_E} > Compras / Entrada de Mercadorias"),
    "w480": ("Compras Mês a Mês - Exportação", f"{_E} > Compras / Entrada de Mercadorias"),
    "v520": ("Pré-Compra", f"{_E} > Compras / Entrada de Mercadorias"), "v540": ("Ultimas Compras", f"{_E} > Compras / Entrada de Mercadorias"),
    "v560": ("Entradas de Mercadorias Registradas", f"{_E} > Compras / Entrada de Mercadorias"),
    "v580": ("Curva ABC por Fornecedor", f"{_E} > Compras / Entrada de Mercadorias"),
    "v600": ("Curva ABC por Produto", f"{_E} > Compras / Entrada de Mercadorias"),

    "w010": ("Cirurgias", f"{_G} > Cirurgias"), "w020": ("Cirurgias Com Valor", f"{_G} > Cirurgias"),
    "w030": ("Cirurgias por Produto", f"{_G} > Cirurgias"), "w040": ("Cirurgias por Data de Faturamento", f"{_G} > Cirurgias"),
    "w050": ("Cirurgias por Instrumentador", f"{_G} > Cirurgias"), "w070": ("Cirurgias com Filtros", f"{_G} > Cirurgias"),
    "w120": ("Cirurgias (Gráfico)", f"{_G} > Cirurgias"), "w125": ("Cirurgias - Materiais Orçados", f"{_G} > Cirurgias"),
    "w130": ("Saídas de Caixas", f"{_G} > Análises de Caixas"), "w140": ("Análise de Saídas de Caixas", f"{_G} > Análises de Caixas"),
    "w150": ("Previsão de Devolução", f"{_G} > Análises de Caixas"), "w160": ("Caixas Perm. - Média de Vendas", f"{_G} > Análises de Caixas"),
    "w180": ("Orçamentos", f"{_G} > Orçamentos"), "w190": ("Orçamentos com Filtros", f"{_G} > Orçamentos"),
    "w200": ("Analítico de Orçamentos", f"{_G} > Orçamentos"), "w210": ("Análise Sintética de Orçamentos", f"{_G} > Orçamentos"),
    "w215": ("Análise de Orçamentos", f"{_G} > Orçamentos"),
    "w220": ("Programação de Entrega", f"{_G} > Pedidos de Venda"), "w230": ("Pedidos (Documentação)", f"{_G} > Pedidos de Venda"),
    "w240": ("Pedidos a Entregar", f"{_G} > Pedidos de Venda"), "w250": ("Pedidos Pendentes", f"{_G} > Pedidos de Venda"),
    "w260": ("Pedidos Cancelados", f"{_G} > Pedidos de Venda"), "w270": ("Pedidos Aguardando Faturamento", f"{_G} > Pedidos de Venda"),
    "w275": ("Pedidos com Quantidades Glosadas", f"{_G} > Pedidos de Venda"),
    "w280": ("Consignações", f"{_G} > Consignações"), "w290": ("Consignações - Exportação p/ Excel", f"{_G} > Consignações"),
    "w300": ("Relatório de Pendências de Consignados", f"{_G} > Consignações"),
    "w310": ("Produtos em Consignação por Período", f"{_G} > Consignações"),
    "w320": ("Materiais Consignados de Terceiros", f"{_G} > Consignações"),
    "w325": ("Programação de Entrega de Consignados", f"{_G} > Consignações"),
    "w330": ("Pedidos de Venda - Exportação", f"{_G} > Análises de Vendas"), "w335": ("Pedidos de Venda - Exportação II", f"{_G} > Análises de Vendas"),
    "w340": ("Vendas por Período", f"{_G} > Análises de Vendas"), "w350": ("Vendas por Produto - Sintética", f"{_G} > Análises de Vendas"),
    "w360": ("Vendas por Produto (Mês/Ano)", f"{_G} > Análises de Vendas"),
    "w370": ("Última Venda de Clientes (com Pedido)", f"{_G} > Análises de Vendas"),
    "w380": ("Últimas Vendas por Cliente", f"{_G} > Análises de Vendas"), "w390": ("Mercadorias Devolvidas", f"{_G} > Análises de Vendas"),
    "w400": ("Produtos Sem Venda", f"{_G} > Análises de Vendas"), "w410": ("Análise de Vendas por Tipo / Subtipo", f"{_G} > Análises de Vendas"),
    "w420": ("Gráficos - Vendas por Produto", f"{_G} > Análises de Vendas"),
    "w430": ("Análise de Lucratividade de Orçamentos", f"{_G} > Análise de Lucratividade"),
    "w440": ("Análise de Lucratividade de Pedidos de Venda", f"{_G} > Análise de Lucratividade"),
    "w510": ("Relatório de Licitações Ganhas", f"{_G} > Licitações"), "w520": ("Relatório de Licitações Perdidas", f"{_G} > Licitações"),
    "w530": ("Acompanhamento de Licitações", f"{_G} > Licitações"), "w535": ("Relatório de Licitações a Definir", f"{_G} > Licitações"),
    "w536": ("Relatório de Licitações Atendidas", f"{_G} > Licitações"),
    "w540": ("Controle de Saídas Avulsas", f"{_G} > Saídas Avulsas"),
    "w560": ("Relação de Lotes", f"{_G} > Auxiliares"), "w570": ("Registro Histórico do Produto", f"{_G} > Auxiliares"),
    "w580": ("Relatório das Saídas na Expedição", f"{_G} > Auxiliares"), "w610": ("Canhotos Não Recebidos", f"{_G} > Auxiliares"),

    "w625": ("Faturamento Com Filtros", f"{_G2} > Faturamento e Analises"), "w630": ("Faturamento por Tipo de Cirurgia", f"{_G2} > Faturamento e Analises"),
    "w635": ("Faturamento Por Cirurgia / Procedimento", f"{_G2} > Faturamento e Analises"),
    "w645": ("Faturamento por Vendedor / Faixa de Lucro", f"{_G2} > Faturamento e Analises"),
    "w650": ("Faturamento Por Médico", f"{_G2} > Faturamento e Analises"), "w660": ("Faturamento por Lote", f"{_G2} > Faturamento e Analises"),
    "w665": ("Faturamento por NCM", f"{_G2} > Faturamento e Analises"), "w675": ("Faturamento Mês a Mês", f"{_G2} > Faturamento e Analises"),
    "w680": ("Relatório de Faturamento", f"{_G2} > Faturamento e Analises"),
    "w690": ("Faturamento Mês a Mês c/ Gráfico", f"{_G2} > Faturamento e Analises"),
    "w695": ("Faturamento Comparativo Anual c/ Gráfico", f"{_G2} > Faturamento e Analises"),
    "w700": ("Faturamento por Vendedor Mês a Mês c/ Gráfico", f"{_G2} > Faturamento e Analises"),
    "w705": ("Notas Fiscais Emitidas sem Vendedor", f"{_G2} > Faturamento e Analises"),
    "w720": ("Notas Fiscais de Saída e Entrada", f"{_G2} > Faturamento e Analises"),
    "w795": ("Vendas de Produtos Tributados PIS/COFINS", f"{_G2} > PIS/COFINS"), "w800": ("Apuração de PIS/COFINS", f"{_G2} > PIS/COFINS"),
    "w805": ("Apuração de PIS/COFINS Analítica", f"{_G2} > PIS/COFINS"),
    "w810": ("Apuração de PIS/COFINS - Data Recebimento", f"{_G2} > PIS/COFINS"),
    "w815": ("Apuração de PIS/COFINS Analítica - NF Recebidas", f"{_G2} > PIS/COFINS"),
    "w830": ("Faturamento de Serviços com Filtros", f"{_G2} > Faturamento de Serviços"),
    "w835": ("Faturamento de Serviços por Conta do Plano", f"{_G2} > Faturamento de Serviços"),

    "w900": ("Etiquetas - Individuais", f"{_A} > Etiquetas (Papel)"), "w905": ("Etiquetas - Clientes Selecionados", f"{_A} > Etiquetas (Papel)"),
    "w910": ("Etiquetas - Aniversários", f"{_A} > Etiquetas (Papel)"), "w915": ("Etiquetas - Remetente", f"{_A} > Etiquetas (Papel)"),
    "w920": ("Etiquetas - Fornecedores", f"{_A} > Etiquetas (Papel)"), "w925": ("Etiquetas - Para o Correio", f"{_A} > Etiquetas (Papel)"),
    "w930": ("Etiquetas - Produto (Jato)", f"{_A} > Etiquetas (Papel)"),
}

# Nome que a planilha 6699 usa para a tela, quando o nome do arquivo não deixa óbvio
NOME_NA_PLANILHA = {
    "w620": "Faturamento",
    "w490": "Apuração de Lucro por NF",
    "w745": "Exportação de Notas Fiscais de Venda",
    "w865": "Exportação de NF Emitidas",
    "u240": "Produtos por tributação",
    "u160": "Relatório de Produtos com Filtros",
    "v620": "Relatório de Notas Fiscais Recebidas",
    "v670": "Exportação de Notas Fiscais Recebidas",
    "manual-cadastro-produto": "Exportação de produtos para Excel",
    "manual-cadastro-bionexo": "Exportação de produtos para Excel",
    "v040": "Relatório da posição de estoque – Fechamento",
    "w100": "Cirurgias Exportação P/Excel",
    "w470": "Análise de Lucratividade por NF",
    "w790-nf-recebidas": "Apuração de PIS/COFINS Analítica - NF Recebidas",
    "w790-natureza": "Apuração de PIS/COFINS - Com Natureza da Rec.",
}


# Impressões de telas de cadastro (orçamento, pedido, licitação, espelho de NF, nota de débito): não são relatórios, ficam fora do catálogo
EXCLUIR = {"j010", "j100", "j200", "m000", "m790"}
IGNORAR_NA_PLANILHA = {
    "Orçamentos – Cadastro", "Licitações – Cadastro de Licitações", "Pedidos de venda – Cadastro de pedidos de venda",
    "Faturamento de pedido de venda – Preparação para faturamento", "Nota de Debito – Cadastro de Notas de debito",
}

# V7 ----------------------------------------------------------------------------------------------------------------
# Na V7 o código vem no nome do PDF (spica_l300_l_faturamento_cc.pdf) e em parte das exportações (L982_..., SPICA_P120_...).
# Nome e área vêm dos menus do SPICA 7 (prints de 27/10/2026). As exportações sem código foram casadas com o item do
# menu pela ordem em que foram gravadas (o levantamento seguiu a ordem do menu).
_EP = "Estoques > Listagens > Posição do Estoque"
_EC = "Estoques > Listagens > Compras / Entrada de Mercadorias"
_EF = "Estoques > Listagens > Em Falta"
_EA = "Estoques > Listagens > Auxiliares"
_EG = "Estoques > Listagens > Ativo Fixo"
_LB = "L.Operacionais > Básicas"
_LG = "L.Operacionais > Gerenciais"
_FP = "Financeiro > Contas a Pagar > Listagens"
_FR = "Financeiro > Contas a Receber > Listagens"

MENU_V7 = {  # código -> (nome da tela no menu, aba > grupo do menu)
    "e300": ("Posição do Estoque", _EP), "e310": ("Posição do Estoque (Fechamento)", _EP),
    "e312": ("Posição do Estoque Com Previsões", _EP), "e313": ("Posição do Estoque Com Valorização", _EP),
    "e320": ("Posição do Estoque Que Estão Fora", _EP), "e330": ("Produtos em Estoque / Consignados", _EP),
    "e340": ("Inventário Valorizado", _EP), "ods-pos-estoque": ("Posição do Estoque - Exportação", _EP),
    "ods-pos-estoque-almox": ("Posição do Estoque - Exportação - Detalhada", _EP),
    "e352": ("Posição do Estoque por RVS", _EP), "ods-pos-estoque-cfop": ("Posição do Estoque - Exportação (CFOP)", _EP),
    "ods-pos-estoque-fech-ii": ("Posição do Estoque - Exportação (Fechamento II)", _EP),

    "e410": ("Pedidos de Compra", _EC), "e420": ("Compras Sintética", _EC), "e430": ("Compras Sintética - Detalhada", _EC),
    "e440": ("Previsão de Compras", _EC), "e460": ("Últimas Compras", _EC), "e470": ("NF Recebidas", _EC),
    "e480": ("Entradas de Mercadorias Registradas", _EC), "e490": ("Compras Mês a Mês", _EC),
    "e494": ("Curva ABC de Fornecedores", _EC), "e497": ("Curva ABC por Produto", _EC), "e498": ("Requisições", _EC),
    "ods-entradas": ("Exportação de Entradas de Mercadorias", _EC), "e505": ("Sugestão de Compras", _EC),
    "ods-requisicoes": ("Requisições - Exportação", _EC), "e509": ("Exportação de Notas Fiscais Recebidas - Geral", _EC),

    "e500": ("Produtos em Falta", _EF), "e510": ("Produtos em Falta nas Caixas", _EF),
    "e520": ("Produtos em Falta em Pedidos", _EF), "e530": ("Produtos em Falta por Cliente", _EF),
    "e540": ("Pendências com Estoque Disponível", _EF),

    "e600": ("Movimentação do Estoque", _EA), "ods-movimentos-data": ("Movimentação do Estoque - Exportação", _EA),
    "e605": ("Movimentação do Estoque de Etiquetas", _EA), "e610": ("Produtos com Lotes a Vencer", _EA),
    "e620": ("Rastreabilidade", _EA), "ods-rastreabilidade": ("Rastreabilidade (Exportação)", _EA),
    "e626": ("Solicitações Pendentes no Almoxarifado", _EA),
    "e628": ("Movimentação do Estoque de Etiquetas (Fabricante)", _EA),
    "ods-mov-estoque-ii": ("Movimentação do Estoque - Exportação II", _EA),

    "e700": ("Posição do Estoque - Ativo Fixo", _EG), "ods-ativo-fixo-ii": ("Posição do Estoque - Ativo Fixo II", _EG),

    "l010": ("Tabela de Preços", f"{_LB} > Tabela de Preços"),
    "l020": ("Produtos Por Tipo", f"{_LB} > Produtos"), "l025": ("Produtos por Fornecedor", f"{_LB} > Produtos"),
    "l030": ("Produtos Com Controle de Lote", f"{_LB} > Produtos"), "l035": ("Produtos Sem Controle de Lote", f"{_LB} > Produtos"),
    "l040": ("Produtos Sem RVS", f"{_LB} > Produtos"), "l050": ("Produtos Por RVS", f"{_LB} > Produtos"),
    "l051": ("Produtos - TUSS", f"{_LB} > Produtos"), "l052": ("Produtos com Filtros 1", f"{_LB} > Produtos"),
    "l054": ("Produtos Sem Foto", f"{_LB} > Produtos"), "l055": ("Produtos Por Família", f"{_LB} > Produtos"),
    "l056": ("Produtos Sem Família", f"{_LB} > Produtos"), "l058": ("Produtos Por Tributação", f"{_LB} > Produtos"),
    "ods-produtos-cad": ("Exporta Cadastro p/ Excel", f"{_LB} > Produtos"),
    "ods-tabela-precos": ("Exportação de Tabela de Preços", f"{_LB} > Produtos"),
    "l053": ("Produtos Por NCM", f"{_LB} > Produtos"), "ods-rel-produtos": ("Exporta Cadastro p/ Excel II", f"{_LB} > Produtos"),
    "l060": ("Caixas", f"{_LB} > Caixas"), "l070": ("Caixas Vendidas", f"{_LB} > Caixas"),
    "l075": ("Caixas com Valor", f"{_LB} > Caixas"),
    "l080": ("Documentos a Vencer por Fornecedor", f"{_LB} > Outras"), "l081": ("Documentos a Vencer por Cliente", f"{_LB} > Outras"),
    "b180": ("RVS a Vencer", f"{_LB} > Outras"), "l090": ("Clientes Bloqueados", f"{_LB} > Outras"),
    "l095": ("Cirurgias a Faturar", f"{_LB} > Outras"),

    "l100": ("Cirurgias", f"{_LG} > Cirurgias"), "l110": ("Cirurgias Com Valor", f"{_LG} > Cirurgias"),
    "l115": ("Cirurgias por Produto", f"{_LG} > Cirurgias"), "l120": ("Cirurgias por Data de Faturamento", f"{_LG} > Cirurgias"),
    "l130": ("Cirurgias por Instrumentador", f"{_LG} > Cirurgias"), "ods-cirurgias": ("Cirurgias (Excel)", f"{_LG} > Cirurgias"),
    "l142": ("Cirurgias com Filtros", f"{_LG} > Cirurgias"), "ods-cirurgias-ii": ("Cirurgias (Excel) II", f"{_LG} > Cirurgias"),
    "l148": ("Cirurgias com Filtros (Méd./Instr.)", f"{_LG} > Cirurgias"),
    "ods-export-cirurgias": ("Exportação de Cirurgias", f"{_LG} > Cirurgias"),
    "ods-export-cirurgias-ocorr": ("Exportação de Cirurgias que Não Ocorreram", f"{_LG} > Cirurgias"),

    "l150": ("Pedidos Pendentes", f"{_LG} > Pedidos"), "l155": ("Pedidos Cancelados", f"{_LG} > Pedidos"),
    "l160": ("Pedidos Fechados e Não Faturados", f"{_LG} > Pedidos"), "l165": ("Pedidos (Documentação)", f"{_LG} > Pedidos"),
    "l170": ("Pedidos a Entregar", f"{_LG} > Pedidos"), "l180": ("Programação de Entrega", f"{_LG} > Pedidos"),
    "l190": ("Pedidos Fechados e Não Faturados II", f"{_LG} > Pedidos"),

    "l200": ("Controle de Saídas de Caixas", f"{_LG} > Análises de Caixas"), "l210": ("Saídas de Caixas", f"{_LG} > Análises de Caixas"),
    "l212": ("Caixas Perm. - Média de Vendas", f"{_LG} > Análises de Caixas"),
    "l214": ("Previsão de Devolução", f"{_LG} > Análises de Caixas"),
    "l216": ("Exportação de Caixas Modelo I", f"{_LG} > Análises de Caixas"),
    "l220": ("Controle de Saídas Avulsas", f"{_LG} > Análises de Saídas Avulsas"),
    "l230": ("Saídas Avulsas Pendentes", f"{_LG} > Análises de Saídas Avulsas"),
    "ods-saidas-avulsas": ("Saídas Avulsas Pendentes - Exportação", f"{_LG} > Análises de Saídas Avulsas"),

    "l270": ("Consignações", f"{_LG} > Consignações"),
    "ods-pend-consig": ("Relatório de Pendências de Consignados", f"{_LG} > Consignações"),
    "l284": ("Produtos em Consignação por Período", f"{_LG} > Consignações"),
    "ods-pend-consig-ii": ("Relatório de Pendências de Consignados II", f"{_LG} > Consignações"),
    "ods-pend-consig-iii": ("Relatório de Pendências de Consignados III", f"{_LG} > Consignações"),
    "l288": ("Programação de Entrega de Consignados", f"{_LG} > Consignações"),
    "l290": ("Consignações - Devoluções / Vendas", f"{_LG} > Consignações"),
    "l292": ("Consignações Pendentes / Finalizadas", f"{_LG} > Consignações"),

    "l300": ("Faturamento", f"{_LG} > Faturamento"), "l305": ("Faturamento Com Filtros", f"{_LG} > Faturamento"),
    "l310": ("Faturamento Por Cirurgia / Procedimento", f"{_LG} > Faturamento"),
    "l320": ("Faturamento Por Vendedor", f"{_LG} > Faturamento"), "l330": ("Faturamento Por Médico", f"{_LG} > Faturamento"),
    "l400": ("Faturamento Resumido", f"{_LG} > Faturamento"),
    "l410": ("Vendas de Produtos Isentos PIS/COFINS", f"{_LG} > Faturamento"),
    "l420": ("Vendas de Produtos Tributados PIS/COFINS", f"{_LG} > Faturamento"),
    "l430": ("Apuração de PIS/COFINS", f"{_LG} > Faturamento"), "l440": ("Apuração de PIS/COFINS Analítica", f"{_LG} > Faturamento"),
    "l441": ("Apuração de PIS/COFINS - Data Recebimento", f"{_LG} > Faturamento"),
    "l442": ("Apuração de PIS/COFINS Analítica - NF Recebidas", f"{_LG} > Faturamento"),
    "l450": ("Faturamento por Lote", f"{_LG} > Faturamento"), "l460": ("Canhotos Não Recebidos", f"{_LG} > Faturamento"),
    "l470": ("Faturamento - Análise de Lucro Bruto", f"{_LG} > Faturamento"),
    "l471": ("Faturamento - Análise de Lucro Bruto (2)", f"{_LG} > Faturamento"),
    "l480": ("Curva ABC de Vendas por Cliente", f"{_LG} > Faturamento"),
    "l490": ("Curva ABC de Vendas por Produto", f"{_LG} > Faturamento"),
    "l500": ("Faturamento Analítico", f"{_LG} > Faturamento"), "l510": ("Faturamento Mês a Mês", f"{_LG} > Faturamento"),
    "l520": ("Faturamento por NCM", f"{_LG} > Faturamento"), "l540": ("Notas Fiscais de Saída e Entrada", f"{_LG} > Faturamento"),
    "l580": ("Notas Fiscais Canceladas", f"{_LG} > Faturamento"), "l590": ("Análise de Lucro Bruto (3)", f"{_LG} > Faturamento"),
    "l595": ("Faturamento Com Filtros - Custo", f"{_LG} > Faturamento"),

    "l610": ("Vendas por Produto (Mês/Ano)", f"{_LG} > Vendas"), "l620": ("Vendas por Período", f"{_LG} > Vendas"),
    "l630": ("Vendas por Produto - Curva ABC", f"{_LG} > Vendas"), "l640": ("Produtos Sem Venda Desde", f"{_LG} > Vendas"),
    "l670": ("Análise de Lucratividade de Pedidos", f"{_LG} > Vendas"), "l680": ("Última Compra de Clientes", f"{_LG} > Vendas"),
    "l690": ("Mercadorias Devolvidas", f"{_LG} > Vendas"), "l695": ("Últimas Vendas por Cliente", f"{_LG} > Vendas"),
    "l696": ("Vendas por Produto - Sintética", f"{_LG} > Vendas"),
    "l698": ("Vendas por Produto (Diretas / com Cirurgia)", f"{_LG} > Vendas"),
    "l635": ("Vendas por Produto - Curva ABC II", f"{_LG} > Vendas"),

    "l700": ("Check List", f"{_LG} > Auxiliares"), "l710": ("Relação de Lotes", f"{_LG} > Auxiliares"),
    "l720": ("Registro Histórico do Produto", f"{_LG} > Auxiliares"), "l730": ("Solicitações para o Estoque", f"{_LG} > Auxiliares"),
    "l750": ("Etiquetas Impressas", f"{_LG} > Auxiliares"), "l780": ("Orçamentos", f"{_LG} > Auxiliares"),
    "l790": ("Ficha de Rastreabilidade", f"{_LG} > Auxiliares"), "l793": ("Orçamentos com Filtros", f"{_LG} > Auxiliares"),
    "l795": ("Orçamentos (Analítica)", f"{_LG} > Auxiliares"), "ods-pedidos-venda": ("Pedidos de Venda (Analítica)", f"{_LG} > Auxiliares"),
    "ods-orcamentos-dc": ("Orçamentos (Dados Complementares)", f"{_LG} > Auxiliares"),
    "ods-piscof-natureza": ("PIS/COFINS – Com Natureza da Receita", f"{_LG} > Auxiliares"),

    "l800": ("Faturamento Geral ABC - NF", f"{_LG} > Novos Relatórios e Análises"),
    "l810": ("Faturamento por Cirurgia", f"{_LG} > Novos Relatórios e Análises"),
    "l820": ("Faturamento Mês a Mês c/ Gráfico", f"{_LG} > Novos Relatórios e Análises"),
    "l830": ("Faturamento Comparativo Anual c/ Gráfico", f"{_LG} > Novos Relatórios e Análises"),
    "l840": ("Faturamento por Vendedor Mês a Mês c/ Gráfico", f"{_LG} > Novos Relatórios e Análises"),
    "l850": ("Notas Fiscais Emitidas sem Vendedor", f"{_LG} > Novos Relatórios e Análises"),
    "l855": ("Análise Sintética de Orçamentos", f"{_LG} > Novos Relatórios e Análises"),
    "l857": ("Análise de Vendas por Tipo / SubTipo", f"{_LG} > Novos Relatórios e Análises"),
    "l860": ("Notas Fiscais Emitidas por Faturista", f"{_LG} > Novos Relatórios e Análises"),
    "l865": ("Faturamento por Vendedor / Faixa de Lucro", f"{_LG} > Novos Relatórios e Análises"),
    "ods-apur-icms-ipi": ("Apuração de ICMS e IPI", f"{_LG} > Novos Relatórios e Análises"),
    "l875": ("Curva ABC de Faturamento por Fabricante", f"{_LG} > Novos Relatórios e Análises"),

    "l900": ("Gráficos - Cirurgias e Procedimentos", f"{_LG} > Gráficos"), "l910": ("Gráficos - Faturamento", f"{_LG} > Gráficos"),
    "l920": ("Gráficos - Vendas por Produto", f"{_LG} > Gráficos"),

    "ods-fat-segmento": ("Faturamento Mensal / Anual por Segmento por Vendedor", f"{_LG} > Exportação"),
    "ods-fat-anual-cli": ("Listagem Anual de Vendas por Vendedor / Cliente", f"{_LG} > Exportação"),
    "l967": ("Apuração de Lucro por Nota Fiscal", f"{_LG} > Exportação"),
    "ods-nf-receb-emit": ("Exportação de Notas Fiscais Recebidas e Emitidas", f"{_LG} > Exportação"),
    "ods-nf-entrada": ("Exportação de Notas Fiscais de Entrada", f"{_LG} > Exportação"),
    "ods-nf-modelo-i": ("Exportação de Notas Fiscais - Modelo I", f"{_LG} > Exportação"),
    "ods-nf-modelo-ii": ("Exportação de Notas Fiscais - Modelo II", f"{_LG} > Exportação"),
    "ods-nf-modelo-iii": ("Exportação de Notas Fiscais - Modelo III", f"{_LG} > Exportação"),
    "ods-nf-modelo-iv": ("Exportação de Notas Fiscais - Modelo IV", f"{_LG} > Exportação"),
    "ods-nf-modelo-v": ("Exportação de Notas Fiscais - Modelo V", f"{_LG} > Exportação"),
    "ods-cliente-completo": ("Relatório Por Cliente Completo", f"{_LG} > Exportação"),
    "ods-cirurgias-modelo-i": ("Exportação de Cirurgias - Modelo I", f"{_LG} > Exportação"),
    "ods-fat-vendedor": ("Faturamento por Vendedor - Exportação", f"{_LG} > Exportação"),
    "ods-vendas-vendedor": ("Vendas por Vendedor - Exportação", f"{_LG} > Exportação"),
    "ods-apur-lucro-iii": ("Apuração de Lucro por Nota Fiscal III", f"{_LG} > Exportação"),
    "ods-exp-nf-prep": ("Exportação de Notas Fiscais Preparadas por Período", f"{_LG} > Exportação"),
    "ods-nf-modelo-vi": ("Exportação de Notas Fiscais - Modelo VI", f"{_LG} > Exportação"),
    "ods-analise-lucr-nf": ("Análise de Lucratividade de NF de Venda", f"{_LG} > Exportação"),
    "ods-exp-nf-faturadas": ("Exportação de Notas Fiscais Faturadas", f"{_LG} > Exportação"),
    "ods-exp-nf-vencidas": ("Exportação de Notas Fiscais Vencidas", f"{_LG} > Exportação"),
    "ods-exp-nf-devol": ("Exportação de Notas Fiscais Emitidas / Devolvidas", f"{_LG} > Exportação"),
    "ods-nf-modelo-vii": ("Exportação de Notas Fiscais - Modelo VII", f"{_LG} > Exportação"),
    "l982": ("Exportação de Notas Fiscais Emitidas - Geral", f"{_LG} > Exportação"),
    "ods-alertas": ("Exportação Manual de Alertas", f"{_LG} > Exportação"),
    "m010": ("Exportação de Análise de Vendas", f"{_LG} > Exportação"),

    "l970": ("Faturamento de Serviços com Filtros", f"{_LG} > Faturamento de Serviços"),
    "l980": ("Materiais Consignados de Terceiros", f"{_LG} > Materiais Consignados"),
    "ods-notas-debito": ("Exportação de Notas de Débito", f"{_LG} > Notas de Débito"),

    "p100": ("Contas a Pagar Por Vencimento", _FP), "p102": ("Contas a Pagar por Emissão", _FP),
    "p110": ("Contas a Pagar por Fornecedor", _FP), "p112": ("Contas a Pagar por Conta do Plano", _FP),
    "p120": ("Contas Pagas", _FP), "p130": ("Contas Pagas Por Fornecedor", _FP), "p135": ("Contas Pagas Por Forma", _FP),
    "p136": ("Contas a Pagar / Pagas Por Forma", _FP), "p140": ("Despesas por Conta do Plano", _FP),
    "p145": ("Despesas por Conta do Plano Mês a Mês", _FP), "p150": ("Comissões", _FP),
    "p160": ("Relatório Mensal de Pagamento de Comissão", _FP), "p170": ("Contas Pagas por Fornecedor / Emissão", _FP),
    "ods-contas-pagar": ("Contas a Pagar - Exportação p/ Excel", _FP), "p190": ("Comissões com Itens", _FP),
    "ods-pagar-pagas": ("Contas a Pagar / Pagas - Exportação", _FP), "ods-comissoes": ("Comissões de Venda - Exportação", _FP),
    "p230": ("Contas à Pagar (Visão em Data Retroativa)", _FP),

    "r100": ("Contas a Receber por Vencimento", _FR), "r102": ("Contas a Receber por Emissão", _FR),
    "r105": ("Contas a Receber por Vendedor", _FR), "r110": ("Contas a Receber por Posição", _FR),
    "r120": ("Contas a Receber por Cliente", _FR), "r130": ("Contas a Receber em Atraso", _FR),
    "r135": ("Contas a Receber por Conta do Plano", _FR), "r140": ("Contas Recebidas", _FR),
    "r150": ("Contas Recebidas por Cliente", _FR), "r155": ("Contas Recebidas por Forma", _FR),
    "r156": ("Contas Recebidas Por NF", _FR), "r158": ("Listagem de Contas Recebidas Por NF Liquidada", _FR),
    "r160": ("Ficha Financeira", _FR), "r170": ("Contas a Receber / Recebidas por No.", _FR),
    "r180": ("Ficha de Duplicatas", _FR), "r190": ("Clientes em Atraso Mês a Mês", _FR),
    "r200": ("Curva ABC Inadimplência", _FR), "r201": ("Relatório de Inadimplência", _FR),
    "ods-receber-recebidas": ("Contas a Receber / Recebidas - Exportação", _FR),
    "r230": ("Contas a Receber / Recebidas - Modelo II", _FR),
    "r240": ("Contas à Receber (Visão em Data Retroativa)", _FR), "ods-receber-vendas": ("Contas à Receber (Vendas)", _FR),
}

# Exportações da V7 sem código no nome: o nome do arquivo sem a data/hora -> código da tela no MENU_V7
ODS_V7 = {
    "Alerta_Notas_Fiscais_de_Devol._Consig._Sinief_Implantes": "ods-alertas",
    "Alerta_Notas_Fiscais_de_Devolução_para_Consignação": "ods-alertas",
    "Analise_Vendas_Tipo_SubTipo": "l857", "Apur_Lucro_NF": "l967", "Apur_Lucro_NF_III": "ods-apur-lucro-iii",
    "Cirurgias": "ods-cirurgias", "Exp_L_Cirurgias_Filtros": "l142", "List_Cirurgias_II": "ods-cirurgias-ii",
    "List_Cirurgias_Filtros_Med_Instr": "l148", "Export_Cirurgias": "ods-export-cirurgias",
    "Export_Cirurgias_N_Ocorr": "ods-export-cirurgias-ocorr", "List_Cirurgias": "ods-cirurgias-modelo-i",
    "Clientes_Atraso_Mes_a_Mes": "r190", "Compras_Mes_a_Mes": "e490", "Despesas_Cta_Plano_Mes_a_Mes": "p145",
    "Exp_NF_Devol": "ods-exp-nf-devol", "Exp_NF_Faturadas": "ods-exp-nf-faturadas", "Exp_NF_Prep_Por_Periodo": "ods-exp-nf-prep",
    "Exp_NF_Vencidas": "ods-exp-nf-vencidas",
    "Export_NF_Modelo_I_Layout_1": "ods-nf-modelo-i", "Export_NF_Modelo_I_Layout_2": "ods-nf-modelo-i",
    "Export_NF_Modelo_I_Layout_3": "ods-nf-modelo-i", "Export_NF_Modelo_I_Layout_4": "ods-nf-modelo-i",
    "Export_NF_Modelo_I_Layout_5": "ods-nf-modelo-i", "Export_NF_Modelo_I_Layout_6": "ods-nf-modelo-i",
    "Export_NF_Modelo_II": "ods-nf-modelo-ii", "Export_NF_Modelo_III": "ods-nf-modelo-iii", "Export_NF_Modelo_IV": "ods-nf-modelo-iv",
    "Export_NF_Modelo_V": "ods-nf-modelo-v", "Export_NF_Modelo_VI": "ods-nf-modelo-vi", "Export_NF_Modelo_VII": "ods-nf-modelo-vii",
    "Faturamento_Analitico": "l500", "Faturamento_Anual_Cli": "ods-fat-anual-cli", "Faturamento_Comparativo": "l830",
    "Faturamento_Geral_ABC": "l800", "Faturamento_Mes_a_Mes": "l820", "Faturamento_Por_Lote": "l450",
    "Faturamento_Tipo_Cir": "l810", "Faturamento_Vend_Mes_a_Mes": "l840", "List_Fat_Vend_Faixa_Lucro": "l865",
    "List_Analise_Lucr_NF_Venda": "ods-analise-lucr-nf", "List_Apur_ICMS_IPI": "ods-apur-icms-ipi", "List_Apur_PISCOF": "l441",
    "List_Pis_Cofins_Nat_Receita": "ods-piscof-natureza",
    "List_C_Pagar_Venc_Fluxo": "p230", "List_C_Receber_Venc_Fluxo": "r240", "List_C_Receber_Recebidas_Nro": "r170",
    "List_C_Recebidas_Forma": "r155", "List_C_Recebidas_NF_Liq": "r158", "List_C_Receber_Vendas": "ods-receber-vendas",
    "List_Comissoes_Vendas": "ods-comissoes", "List_Contas_A_Pagar": "ods-contas-pagar", "List_Contas_Pagar_e_Pagas": "ods-pagar-pagas",
    "List_Contas_Receber_e_Recebidas": "ods-receber-recebidas",
    "List_Entr_Merc": "ods-entradas", "Mercadorias_Reg": "e480", "List_Req_Compras": "ods-requisicoes",
    "List_Faturamento_Vendedor": "ods-fat-vendedor", "List_Vendas_Vendedor": "ods-vendas-vendedor",
    "List_Inv_Val": "e340", "List_Mov_Estoq_Exp_II": "ods-mov-estoque-ii", "List_Movimentos_Data": "ods-movimentos-data",
    "List_Prod_Lotes_a_Venc": "e610", "List_Rastreabilidade": "ods-rastreabilidade", "Prod_Estoq_Consig": "e330",
    "List_Pos_Estoq": "ods-pos-estoque", "List_Pos_Estoq_Det_Almox": "ods-pos-estoque-almox",
    "List_Pos_Estoq_Exp_CFOP": "ods-pos-estoque-cfop", "List_Pos_Estoq_Fech_II": "ods-pos-estoque-fech-ii",
    "List_Pos_Estoque_Ativo_Fixo_II": "ods-ativo-fixo-ii",
    "List_Orcamentos_DC": "ods-orcamentos-dc", "Orcamentos": "l795", "Ped_V": "ods-pedidos-venda", "List_Pedidos_Doc": "l165",
    "List_Produtos_Cad": "ods-produtos-cad", "Rel_Produtos": "ods-rel-produtos", "Tabela_Precos": "ods-tabela-precos",
    "NF_Entrada": "ods-nf-entrada", "NF_Receb_Emit": "ods-nf-receb-emit",
    "Notas_Debito_Cliente": "ods-notas-debito", "Notas_Debito_Medico": "ods-notas-debito", "Notas_Debito_Vendedor": "ods-notas-debito",
    "Produtos_Consig": "l284", "Rel_Pend_Consig": "ods-pend-consig", "Rel_Pend_Consig_II": "ods-pend-consig-ii",
    "Rel_Pend_Consig_III": "ods-pend-consig-iii", "Saidas_Avulsas_Pend": "ods-saidas-avulsas",
    "Rel_Cliente_Completo": "ods-cliente-completo", "Rel_Faturamento_Fab": "ods-fat-segmento", "Rel_Faturamento_Tipo": "ods-fat-segmento",
}


NOME_NA_PLANILHA_V7 = {  # nome que a aba SPICA7 da planilha 6699 usa, quando difere do menu
    "e340": "Posição do Estoque – Inventário Valorizado", "e430": "Pedidos de Compra – Sintética Detalhada",
    "e470": "Notas Fiscais Recebidas", "ods-entradas": "Exportação de Entradas", "e509": "Exportação de NF Recebidas – Geral",
    "ods-ativo-fixo-ii": "Ativo Fixo II", "l300": "Listagem do Faturamento", "l590": "Faturamento – Análise de Lucro Bruto (3)",
    "l967": "Apuração de Lucro por NF", "ods-apur-lucro-iii": "Apuração de Lucro por NF (III)", "l970": "Faturamento de Serviços",
    "ods-exp-nf-devol": "Exportação de NF Devolvidas", "l982": "Exportação de NF Emitidas – Geral",
    "ods-piscof-natureza": "Apuração de PIS/COFINS - Natureza da Receita", "ods-nf-receb-emit": "Exportação de NF Emitidas e Recebidas",
}


def base_do_ods(nome):
    """Nome da exportação sem a data/hora em que foi gerada: List_Inv_Val_20261021_022005_2.ods -> List_Inv_Val."""
    return re.sub(r"_(?:\d{8}|20\d\d)_\d{6}(?:_\d+)?$", "", Path(nome).stem)


def classificar_v7(arquivo):
    """Código da tela do arquivo da V7: o do nome (spica_l300_..., L982_...) ou o da tabela ODS_V7."""
    nome = arquivo["nome"]
    achou = re.match(r"^(?:spica_)?([a-z]\d{3})_", nome, re.I)
    if achou:
        return achou.group(1).lower()
    return ODS_V7.get(base_do_ods(nome)) or "ods-" + sem_acento(base_do_ods(nome)).replace("_", "-")


def nome_pelo_titulo(titulos):
    """Nome da tela a partir dos títulos dos arquivos: o mais comum, sem 'Listagem de' e sem o 'entre' do período."""
    if not titulos:
        return None
    titulo = max(set(titulos), key=lambda t: (titulos.count(t), -len(t)))
    titulo = re.sub(r"\s+(?:entre|desde|Entre)$", "", titulo)
    titulo = re.sub(r"^(?:Listagem|Relat[óo]rio)\s+(?:de|do|da|dos|das)?\s*", "", titulo, flags=re.I).strip(" -")
    return titulo[:1].upper() + titulo[1:] if titulo else None


def sem_acento(texto):
    return "".join(c for c in unicodedata.normalize("NFKD", texto.lower()) if not unicodedata.combining(c))


def codigo_do_nome(nome):
    achou = re.match(r"^([a-z]\d{3})", nome, re.I)
    return achou.group(1).lower() if achou else None


def classificar(arquivo):
    """Rotina a que o arquivo pertence: o código do nome, ou uma regra para os que não têm código."""
    nome = arquivo["nome"]
    codigo = codigo_do_nome(nome)
    if codigo == "w790":  # o nome dos 7 PDFs é igual (..._isentos_piscof_N); quem separa as rotinas é o título
        titulo = sem_acento(arquivo["titulo"])
        for chave, rotina in (("nf recebidas", "w790-nf-recebidas"), ("natureza", "w790-natureza"), ("isentos", "w790-isentos"),
                              ("trib", "w790-tributados"), ("data de recebimento", "w790-recebimento"), ("analitica", "w790-analitica")):
            if chave in titulo:
                return rotina
        return "w790-apuracao"
    if codigo:
        return codigo
    base = sem_acento(nome)
    for prefixo, rotina in (
        ("mapa_um_mes", "x450"), ("mapa_despesas_receitas_comparativo", "x510"), ("mapa_despesas_receitas_compet", "x490"),
        ("mapa_despesas_receitas_prev_real", "x520"), ("mapa_despesas_receitas", "x480"),
        ("nf_compra", "manual-nf-compra"), ("posicao_estoque_exporta", "manual-posicao-estoque"),
        ("rel_exp_cad_bionexo", "manual-cadastro-bionexo"), ("rel_exp_cad_produto", "manual-cadastro-produto"),
        ("list_pos_estoq_det_almox", "ods-pos-estoque-almox"), ("list_pos_estoq_det_consig", "ods-pos-estoque-consig"),
        ("list_pos_estoq", "ods-pos-estoque"), ("movimentacao_estoq_export", "ods-mov-estoque"),
        ("rastreabilidade_exp", "ods-rastreabilidade"), ("rel_exp_tab_pr", "ods-tabela-precos"),
    ):
        if base.startswith(prefixo):
            return rotina
    return "ods-" + re.sub(r"_?\d{8}_\d{6}.*$", "", base).replace("_", "-")  # exportação nova que ainda não tem regra: vira tela própria


def titulo_do_arquivo(caminho):
    """Título do relatório: primeira linha da exportação, ou o 'Relatório de...' no texto do PDF."""
    if caminho.suffix.lower() == ".ods":
        xml = zipfile.ZipFile(caminho).read("content.xml").decode("utf8")
        for linha in re.findall(r"<table:table-row[^>]*>(.*?)</table:table-row>", xml, re.S)[:6]:
            celulas = re.findall(r"<table:table-cell[^>]*?(?:/>|>(.*?)</table:table-cell>)", linha, re.S)
            textos = [" ".join(html.unescape(re.sub(r"<[^>]+>", " ", c)).split()) for c in celulas]
            textos = [t for t in textos if t]
            if textos:  # linha com mais de 3 células já são os nomes das colunas: a exportação não tem título
                return "" if len(textos) > 3 else textos[0][:110]
        return ""
    texto = texto_do_pdf(caminho, max_paginas=2)  # o mesmo extrator do conteúdo, para o título casar com o texto
    # "Fluxo" não entra: é nome de coluna em vários PDFs e aparece antes do título ("Fluxo Fornecedor")
    achou = re.search(r"((?:Relat[óo]rio|Listagem|Curva ABC|Extrato|Gr[áa]fico|Mapa|Ficha|Espelho|Saldos|D\.R\.E\.|Análise|Exporta|R\. [A-Z])[^\n]{0,110})", texto)
    if not achou:
        return ""
    corte = re.split(r"\s\d{2}/\d{2}/\d{4}|\sPag\.|Data:|\s\d{4}\s*[A-Z]|\s\d{2}/\d{4}", achou.group(1))[0]
    return " ".join(corte.split())[:100]


def ler_arquivos(versao):
    achados = {}
    for pasta in PASTAS[versao]:
        if not (RAIZ / pasta).is_dir():
            continue
        for caminho in sorted((RAIZ / pasta).iterdir()):
            if caminho.suffix.lower() not in (".pdf", ".ods") or caminho.name.startswith("6699_Levantamento"):  # a planilha de controle não é relatório
                continue
            arq = achados.setdefault(caminho.name, {"nome": caminho.name, "tipo": caminho.suffix.lower()[1:], "pastas": [], "caminho": caminho})
            arq["pastas"].append(pasta)
    for arq in achados.values():
        arq["titulo_bruto"] = titulo_do_arquivo(arq["caminho"])
        if re.match(r"Joinner\b", arq["titulo_bruto"]):  # cabeçalho com empresa e usuário, não é o título
            arq["titulo_bruto"] = ""
        arq["titulo"] = limpar_titulo(arq["titulo_bruto"])  # sem o período: as datas só dependem do dia em que foi exportado
        arq["cita_imposto"] = versao in arq["pastas"]
        arq["rotina"] = classificar_v7(arq) if versao == "V7" else classificar(arq)
    return list(achados.values())


def nomes_da_planilha(versao):
    arquivo = PLANILHA if PLANILHA.exists() else PLANILHA.with_name(PLANILHA.name + ".bak")  # o SPICA/LibreOffice deixa só o .bak se a planilha for apagada
    quadro = pd.read_excel(arquivo, sheet_name=ABA_PLANILHA[versao], engine="odf", header=None)
    return [str(v).strip() for v in quadro[0].dropna() if str(v).strip().lower() != "rotina"]


def nome_e_area(versao, rotina, arqs):
    """(nome, área) da tela; nome None quando não dá para saber."""
    if versao == "V8":
        if rotina in ROTINAS:
            return ROTINAS[rotina]
        if rotina in MENU_V8:
            return MENU_V8[rotina]
        return nome_pelo_titulo([a["titulo"] for a in arqs if a["titulo"]]), "SEM CLASSIFICAÇÃO"
    if rotina in MENU_V7:
        return MENU_V7[rotina]
    return nome_pelo_titulo([a["titulo"] for a in arqs if a["titulo"]]), "SEM CLASSIFICAÇÃO"


def so_letras(texto):
    return " ".join(re.sub(r"[^a-z0-9]+", " ", sem_acento(texto)).split())


def melhor_nome_da_planilha(nome_proposto, planilha):
    # ratio e não token_set_ratio: este dá 100% quando um nome está contido no outro ("Faturamento" x "Gráficos de Faturamento")
    pontos = [(fuzz.ratio(so_letras(nome_proposto), so_letras(p)), p) for p in planilha]
    return max(pontos) if pontos else (0, "")


def montar(versao):
    arquivos = [a for a in ler_arquivos(versao) if a["rotina"] not in EXCLUIR]
    planilha = [p for p in nomes_da_planilha(versao) if p not in IGNORAR_NA_PLANILHA]
    por_rotina = defaultdict(list)
    for arq in arquivos:
        por_rotina[arq["rotina"]].append(arq)

    linhas = []
    for rotina, arqs in por_rotina.items():
        proposto, area = nome_e_area(versao, rotina, arqs)
        explicito = (NOME_NA_PLANILHA if versao == "V8" else NOME_NA_PLANILHA_V7).get(rotina)
        nota, melhor = (100, explicito) if explicito else melhor_nome_da_planilha(proposto or "", planilha)
        linhas.append({
            "versao": versao, "rotina": rotina, "nome": proposto or "(revisar)", "area": area,
            "pdfs": sum(a["tipo"] == "pdf" for a in arqs), "exportacoes": sum(a["tipo"] == "ods" for a in arqs),
            "imposto": sum(a["cita_imposto"] for a in arqs), "arquivos": arqs,
            "planilha": melhor if nota >= 88 else "", "nota_planilha": nota,
            "titulos": sorted({a["titulo"] for a in arqs if a["titulo"]}),
            "parcial": all(a["pastas"] == [versao] for a in arqs),  # só os que citam imposto foram salvos
        })
    linhas.sort(key=lambda l: (l["area"], l["rotina"]))
    usados = {l["planilha"] for l in linhas if l["planilha"]}
    return linhas, [p for p in planilha if p not in usados]


def salvar(linhas, sem_arquivo, versao):
    wb = Workbook()
    aba = wb.active
    aba.title = "Rotinas"
    titulos = ["Código", "Nome da rotina (corrija aqui)", "Área (palpite)", "PDFs", "Exportações", "Cita imposto",
               "Cobertura", "Nome na planilha 6699", "Títulos vistos nos arquivos"]
    aba.append(titulos)
    for l in linhas:
        cobertura = "só os que citam imposto" if l["parcial"] else "completa"
        aba.append([l["rotina"], l["nome"], l["area"], l["pdfs"], l["exportacoes"], l["imposto"] or "", cobertura,
                    l["planilha"], " | ".join(l["titulos"])[:300]])
    for c in aba[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2B579A")
    for coluna, largura in zip("ABCDEFGHI", (16, 52, 30, 7, 12, 12, 24, 44, 90)):
        aba.column_dimensions[coluna].width = largura
    aba.freeze_panes = "C2"
    aba.auto_filter.ref = aba.dimensions

    arq_aba = wb.create_sheet("Arquivos")
    arq_aba.append(["Código", "Rotina", "Arquivo", "Tipo", "Cita imposto", "Pastas", "Título dentro do arquivo"])
    for l in linhas:
        for a in l["arquivos"]:
            arq_aba.append([l["rotina"], l["nome"], a["nome"], a["tipo"], "sim" if a["cita_imposto"] else "", " + ".join(a["pastas"]), a["titulo"]])
    for c in arq_aba[1]:
        c.font = Font(bold=True)
    for coluna, largura in zip("ABCDEFG", (16, 48, 62, 7, 12, 18, 80)):
        arq_aba.column_dimensions[coluna].width = largura

    pla = wb.create_sheet("Planilha sem arquivo")
    pla.append([f"Rotinas da planilha 6699 ({ABA_PLANILHA[versao]}) que não bateram com nenhuma rotina do catálogo"])
    pla["A1"].font = Font(bold=True)
    for nome in sem_arquivo:
        pla.append([nome])
    pla.column_dimensions["A"].width = 90
    for linha in aba.iter_rows(min_row=2):
        for c in linha:
            c.alignment = Alignment(vertical="top", wrap_text=c.column_letter in "BHI")
    wb.save(saida(versao))


def texto_do_pdf(caminho, max_paginas=None):
    try:
        with pymupdf.open(caminho) as doc:
            return "\n".join(pagina.get_text() for pagina in list(doc)[:max_paginas])
    except Exception:
        return "\n".join((p.extract_text() or "") for p in pypdf.PdfReader(str(caminho)).pages[:max_paginas])


def linhas_da_ods(caminho):
    """Linhas não vazias da exportação, todas as abas, cada uma como lista de células."""
    xml = zipfile.ZipFile(caminho).read("content.xml").decode("utf8")
    linhas = []
    for linha in re.findall(r"<table:table-row[^>]*>(.*?)</table:table-row>", xml, re.S):
        celulas = re.findall(r"<table:table-cell[^>]*?(?:/>|>(.*?)</table:table-cell>)", linha, re.S)
        textos = [" ".join(html.unescape(re.sub(r"<[^>]+>", " ", c)).split()) for c in celulas]
        textos = [t for t in textos if t]
        if textos:
            linhas.append(textos)
    return linhas


def conteudo_do_arquivo(arq):
    """(colunas, texto), os dois sem o título (ele fica em arq["titulo"] e é mostrado como nome do layout).
    Colunas: os nomes das colunas da exportação ou o cabeçalho do relatório em PDF. Texto: o resto do arquivo."""
    if arq["tipo"] == "ods":
        linhas = linhas_da_ods(arq["caminho"])
        if not linhas:
            return "", ""
        corpo = linhas[1:] if len(linhas[0]) <= 3 else linhas  # primeira linha curta = título; longa = já são as colunas
        # Os nomes das colunas são as linhas do começo que não têm número (o cabeçalho pode ter mais de uma linha).
        cabecalho = []
        for linha in corpo[:4]:
            if len(linha) >= 2 and sum(bool(re.fullmatch(r"[\d.,%/\s-]+", c)) for c in linha) / len(linha) < 0.25:
                cabecalho.append(linha)
            else:
                break
        if not cabecalho and corpo:  # sem linha só de texto: a mais larga do começo
            cabecalho = [max(corpo[:8], key=len)]
        colunas = " | ".join(c for linha in cabecalho for c in linha) if cabecalho and len(cabecalho[0]) >= 2 else ""
        texto = "\n".join(" | ".join(l) for l in corpo)
    else:
        texto = texto_do_pdf(arq["caminho"])
        bruto = arq.get("titulo_bruto", "")
        if bruto and bruto in texto:
            texto = texto.replace(bruto, " ", 1)
        colunas = " ".join(texto[:3000].split())  # onde ficam os nomes dos campos do relatório
    return colunas, texto[:LIMITE_TEXTO]


def salvar_banco(linhas):
    """Grava o catálogo para o bot consultar. Escreve num arquivo novo e troca no fim, para não deixar o banco pela metade."""
    novo = BANCO.with_suffix(".db.novo")
    novo.unlink(missing_ok=True)
    con = sqlite3.connect(novo)
    con.executescript("""
        CREATE TABLE rotinas (codigo TEXT, versao TEXT, nome TEXT, area TEXT, parcial INTEGER, cita_imposto INTEGER,
                              PRIMARY KEY (versao, codigo));
        CREATE TABLE arquivos (id INTEGER PRIMARY KEY, codigo TEXT, versao TEXT, nome_arquivo TEXT, tipo TEXT,
                               caminho TEXT, cita_imposto INTEGER, titulo TEXT, colunas TEXT, texto TEXT);
    """)
    total = 0
    for l in linhas:
        con.execute("INSERT INTO rotinas VALUES (?, ?, ?, ?, ?, ?)",
                    (l["rotina"], l["versao"], l["nome"], l["area"], int(l["parcial"]), int(bool(l["imposto"]))))
        for arq in l["arquivos"]:
            colunas, texto = conteudo_do_arquivo(arq)
            con.execute("INSERT INTO arquivos (codigo, versao, nome_arquivo, tipo, caminho, cita_imposto, titulo, colunas, texto) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
                        (l["rotina"], l["versao"], arq["nome"], "pdf" if arq["tipo"] == "pdf" else "exportacao", str(arq["caminho"]),
                         int(arq["cita_imposto"]), arq["titulo"], colunas, texto))
            total += 1
    con.commit()
    con.close()
    os.replace(novo, BANCO)
    return total


def reconstruir():
    """Relê as pastas das versões, regera as planilhas de revisão e o banco do bot. Retorna (rotinas, arquivos)."""
    todas = []
    for versao in VERSOES:
        linhas, sem_arquivo = montar(versao)
        salvar(linhas, sem_arquivo, versao)
        todas += linhas
    return len(todas), salvar_banco(todas)


if __name__ == "__main__":
    todas = []
    for versao in VERSOES:
        linhas, sem_arquivo = montar(versao)
        salvar(linhas, sem_arquivo, versao)
        todas += linhas
        print(f"{versao}: {len(linhas)} rotinas, {sum(len(l['arquivos']) for l in linhas)} arquivos -> {saida(versao).name}")
        print("  sem nome:", [l["rotina"] for l in linhas if l["nome"] == "(revisar)"])
        print("  planilha sem arquivo:", sem_arquivo)
    print(f"{salvar_banco(todas)} arquivos -> {BANCO.name}")
